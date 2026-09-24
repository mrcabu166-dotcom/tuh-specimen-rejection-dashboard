"""Regression checks for the monthly Google Sheets workflow."""

import io
import unittest

from openpyxl import Workbook

from cleaner import load_and_consolidate, parse_sheet_month


HEADERS = ['วันที่', 'เวลา', 'HN', 'Ward', 'สิ่งส่งตรวจ', 'ใบส่งตรวจ']


def make_workbook(november_has_case=False, with_denominator=False):
    book = Workbook()
    book.remove(book.active)

    def monthly_sheet(name, day=None, header=True):
        sheet = book.create_sheet(name)
        if header:
            sheet.append(HEADERS)
        if day is not None:
            sheet.append([day, '09:00', 'HN-ปิดบัง', 'อช.ส', 'ติดชื่อผิดราย', None])

    # Staff may drag tabs, so workbook order must not determine chronology.
    monthly_sheet('ม.ค. 70', 2)
    monthly_sheet('พ.ย. 69', 3 if november_has_case else None)
    monthly_sheet('ต.ค. 69', 1)
    monthly_sheet('ธ.ค. 69', 5, header=False)  # incomplete new tab
    monthly_sheet('ต.ค.', 4)  # missing year, cannot be assigned safely
    monthly_sheet('หมายเหตุ ต.ค. 69', 6)  # helper tab, not a monthly tab
    monthly_sheet('สรุป', 7)  # existing summary, must not double count
    if with_denominator:
        totals = book.create_sheet('ยอดตรวจทั้งหมด')
        totals.append(['เดือน', 'จำนวนสิ่งส่งตรวจทั้งหมด'])
        totals.append(['ต.ค. 69', 100])
        totals.append(['พ.ย. 69', 200])
        totals.append(['ม.ค. 70', 300])

    result = io.BytesIO()
    book.save(result)
    result.seek(0)
    result.name = 'monthly.xlsx'
    return result


class MonthlyTabsTest(unittest.TestCase):
    def test_month_names_require_unambiguous_year(self):
        self.assertEqual(parse_sheet_month('ต.ค. 69')[:2], (10, 2026))
        self.assertEqual(parse_sheet_month('ม.ค.70')[:2], (1, 2027))
        self.assertEqual(parse_sheet_month('ตุลาคม 2569')[:2], (10, 2026))
        self.assertEqual(parse_sheet_month('ต.ค. 2026')[:2], (10, 2026))
        self.assertIsNone(parse_sheet_month('ต.ค.')[0])
        self.assertIsNone(parse_sheet_month('หมายเหตุ ต.ค. 69')[0])

    def test_new_month_is_read_and_empty_future_tab_is_safe(self):
        cases, causes = load_and_consolidate(make_workbook())
        self.assertEqual(cases['date'].tolist(), ['2026-10-01', '2027-01-02'])
        self.assertEqual(cases['fiscal_year'].tolist(), ['ปีงบประมาณ 2570'] * 2)
        self.assertEqual(len(causes), 2)
        self.assertTrue(any('ธ.ค. 69' in warning for warning in cases.attrs['ingestion_warnings']))
        self.assertTrue(any('ชื่อไม่ครบ' in warning for warning in cases.attrs['ingestion_warnings']))

    def test_rows_added_later_appear_without_code_change(self):
        cases, _ = load_and_consolidate(make_workbook(november_has_case=True))
        self.assertEqual(cases['date'].tolist(), ['2026-10-01', '2026-11-03', '2027-01-02'])

    def test_yearless_only_tab_is_not_assigned_to_wrong_fiscal_year(self):
        book = Workbook()
        sheet = book.active
        sheet.title = 'ต.ค.'
        sheet.append(HEADERS)
        sheet.append([1, '09:00', 'HN-ปิดบัง', 'อช.ส', 'ติดชื่อผิดราย', None])
        result = io.BytesIO()
        book.save(result)
        result.seek(0)
        result.name = 'yearless.xlsx'

        cases, _ = load_and_consolidate(result)
        self.assertTrue(cases.empty)
        self.assertTrue(cases.attrs['ingestion_warnings'])

    def test_optional_denominator_sheet_is_loaded(self):
        cases, _ = load_and_consolidate(make_workbook(with_denominator=True))
        denominator = cases.attrs['denominator']
        self.assertEqual(denominator['year_month'].tolist(), ['2026-10', '2026-11', '2027-01'])
        self.assertEqual(denominator['total_specimens'].tolist(), [100.0, 200.0, 300.0])
        self.assertTrue(cases.attrs['denominator_sheet_found'])


if __name__ == '__main__':
    unittest.main()
